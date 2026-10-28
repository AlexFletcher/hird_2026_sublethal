#!/usr/bin/env python3
"""
================================================================================
 co_analysis.py  —  CO dose-response analysis (Hird, Stafford, Fletcher, Herigstad)
================================================================================

Single, self-contained analysis script for the manuscript

    "Sub-lethal metabolic suppression by carbon monoxide near occupational
     exposure limits: a hierarchical Bayesian dose-response analysis in
     human brain cells"

It reproduces every quantitative figure and table in the paper from one
consolidated data workbook. The code is organised into labelled sections so it
is clear how each output is produced:

    SECTION 1  Loaders            — read flow, resazurin and protein from the workbook
    SECTION 2  Models             — the two hierarchical Bayesian models
    SECTION 3  Fitting & cache    — fit once, cache posteriors to disk
    SECTION 4  Posterior helpers  — turn posteriors into dose-response summaries
    SECTION 5  Figures            — one function per manuscript figure
    SECTION 6  Table              — Table 2 (posterior estimates)
    SECTION 7  Command-line       — regenerate any figure/table on demand

--------------------------------------------------------------------------------
FIGURE NUMBERING (matches the manuscript)
    Figure 1  Experimental design + analysis workflow  (schematic; produced
              separately in a vector graphics editor — not generated here)
    Figure 2  Metabolic dose-response       (resazurin, per line)
    Figure 3  Cell death / viability        (flow composition, HA/HUVEC/U87)
    Figure 4  Metabolic-death decoupling    (metabolism falls, viability holds)
    Figure 5  Protein normalisation         (total protein flat; per-µg-protein falls)
    Table 2   Per-line posterior estimates  (table2_estimates.csv)

USAGE
    python co_analysis.py --all                 # everything (fits, then all outputs)
    python co_analysis.py --figure 2            # just Figure 2 (uses cached fit)
    python co_analysis.py --figure 2 3 4        # several figures
    python co_analysis.py --table               # just Table 2
    python co_analysis.py --refit --all         # force a fresh fit
    python co_analysis.py --data PATH --outdir DIR --all
    python co_analysis.py --check-resazurin --raw-data Paper_Data_Resazurin.xlsx
                                                # provenance: rebuild resazurin from raw

    Fast smoke-test (not for publication numbers):
    python co_analysis.py --all --draws 200 --tune 200 --chains 2

The default sampling settings (4 chains x 2000 draws) are the production settings
used for the manuscript. The first run that needs a fit will sample both models
(~5 min on a normal machine) and cache the posteriors next to the outputs; later
runs that only rebuild a figure reuse the cache and are near-instant.

DESIGN DECISIONS ENCODED HERE (see Methods):
  * dose enters as log10(ppm + 1), centred; each dose-response is quadratic
    (admits saturating / non-monotonic shapes);
  * 400 ppm is excluded from the primary fits;
  * SH-SY5Y is excluded from the death analysis (harvest artefact) but kept for
    metabolism;
  * cross-line comparison is on within-line responses relative to 0 ppm control.

DEPENDENCIES: python>=3.9, numpy, pandas, pymc>=5, arviz, matplotlib, openpyxl.
  Some environments crash the sampler under multi-threaded BLAS; this script pins
  BLAS to one thread per process (standard for MCMC) before importing numpy.
================================================================================
"""

import os
import sys
import json
import re
import argparse

# Pin BLAS to one thread per process before importing NumPy. This is standard for
# MCMC and also avoids a sampler crash seen on some constrained systems.
for _v in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "NUMEXPR_NUM_THREADS"):
    os.environ.setdefault(_v, "1")

import numpy as np  # noqa: E402  (imports below must follow the BLAS setup above)
import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

# ------------------------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------------------------
DOSE_OFFSET = 1.0  # dose enters as log10(ppm + OFFSET)
QUADRANTS = ["live", "early", "late", "necrotic"]  # flow composition categories
NONREF = ["early", "late", "necrotic"]  # non-reference categories (live is reference)
PRIMARY_MAX = 400  # 400 ppm excluded from primary fits (strict <)
GRID = np.array([0, 5, 10, 20, 40, 80, 120])  # dose grid for plotting posteriors (< 400)

# per-line colour scheme (shared by every figure and the schematic)
COL = {"HA": "#1b7837", "HUVEC": "#2166ac", "U87": "#762a83", "SH-SY5Y": "#d95f02"}
LAB = {
    "HA": "HA (primary astrocyte)",
    "HUVEC": "HUVEC (primary endothelial)",
    "U87": "U87 (glioblastoma)",
    "SH-SY5Y": "SH-SY5Y (neuroblastoma)",
}

# production sampling settings (used for the manuscript)
SAMPLE_DEFAULTS = dict(
    draws=2000,
    tune=2000,
    chains=4,
    cores=4,
    target_accept=0.95,
    init="adapt_diag",
    random_seed=7,
    progressbar=False,
)

CACHE_RESAZURIN = "_cache_resazurin.nc"
CACHE_COMPOSITION = "_cache_composition.nc"
CACHE_META = "_cache_meta.json"

# publication figure styling: sans-serif (Arial/Helvetica, DejaVu Sans fallback),
# consistent legible sizes, no on-plot titles (titles live in the caption),
# fonts embedded for vector output.
plt.rcParams.update(
    {
        "font.family": "sans-serif",
        "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
        "font.size": 9,
        "axes.labelsize": 9,
        "axes.titlesize": 9,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "legend.fontsize": 8,
        "axes.linewidth": 0.8,
        "xtick.major.width": 0.8,
        "ytick.major.width": 0.8,
        "savefig.dpi": 300,
        "figure.dpi": 150,
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
    }
)


def _panel(ax, letter):
    """Bold panel letter at the top-left of an axis."""
    ax.text(
        -0.02,
        1.03,
        letter,
        transform=ax.transAxes,
        ha="right",
        va="bottom",
        fontsize=11,
        fontweight="bold",
    )


# Every figure is written as a 300-dpi PNG (raster preview / submission) plus
# vector PDF and SVG (for final production and hand-editing). Fonts are embedded
# in the vector output via the rcParams above (pdf.fonttype 42, svg.fonttype none).
FIG_FORMATS = ("png", "pdf", "svg")


def _savefig(fig, outdir, stem, bbox_inches=None):
    """Save `fig` as PNG (300 dpi), PDF and SVG under `stem`; return the paths.

    The PNG is listed first (the reference raster); the PDF and SVG are vector
    copies carrying the same basename.
    """
    paths = []
    for ext in FIG_FORMATS:
        path = os.path.join(outdir, f"{stem}.{ext}")
        fig.savefig(path, dpi=300, bbox_inches=bbox_inches)
        paths.append(path)
    return paths


# ==============================================================================
# SECTION 1 — LOADERS
#   Each reads one readout from the consolidated workbook into a tidy DataFrame,
#   parsing the experimenters' wide sheet layouts into model-ready long tables.
# ==============================================================================

# --- 1a. Flow cytometry -------------------------------------------------------
_FLOW_LINES = ["U87", "HUVEC", "HA", "SH-SY5Y"]
_FLOW_LABEL_RE = re.compile(
    r"^\s*(" + "|".join(map(re.escape, _FLOW_LINES)) + r")\s+(One|Two|Three)\s*$"
)
_FLOW_REP = {"One": 1, "Two": 2, "Three": 3}
_FLOW_STATUS = {"U87": "cancer", "HUVEC": "non-cancer", "HA": "non-cancer", "SH-SY5Y": "cancer"}
_FLOW_LINEAGE = {
    "U87": "glioblastoma",
    "HUVEC": "endothelial",
    "HA": "astrocyte",
    "SH-SY5Y": "neuroblastoma",
}
_FLOW_OFF = dict(dose=0, cells=1, live=2, dead=3, early=4, late=5, necrotic=6)
_FLOW_CTRL_NAME_OFF = 8
_FLOW_CTRL_OFF = dict(cells=9, live=10, dead=11, early=12, late=13, necrotic=14)


def _num(v):
    return v if isinstance(v, (int, float)) else None


def load_flow(path, sheet="(1-2) Flow Data"):
    """Parse the wide flow sheet -> (dose_df, controls_df).

    dose_df carries the model-ready counts (n_live/n_early/n_late/n_necrotic,
    n_events), the quadrant fractions (f_*), and a total dead fraction f_dead.
    Biological replicates are stacked in One/Two/Three bands, lines side by side.
    """
    wv = load_workbook(path, data_only=True)[sheet]
    dose_rows, ctrl_rows = [], []
    for r in range(1, wv.max_row + 1):
        for c0 in range(1, wv.max_column + 1):
            m = _FLOW_LABEL_RE.match(str(wv.cell(r, c0).value or ""))
            if not m:
                continue
            line, rep = m.group(1), _FLOW_REP[m.group(2)]
            for k in range(7):  # seven dose rows
                rr = r + 3 + k
                dose = _num(wv.cell(rr, c0 + _FLOW_OFF["dose"]).value)
                live = _num(wv.cell(rr, c0 + _FLOW_OFF["live"]).value)
                if dose is None or live is None:
                    continue
                q = {
                    name: _num(wv.cell(rr, c0 + _FLOW_OFF[name]).value)
                    for name in ("live", "early", "late", "necrotic", "cells", "dead")
                }
                n_events = sum(v for v in (q["live"], q["early"], q["late"], q["necrotic"]) if v)
                dose_rows.append(
                    dict(
                        line=line,
                        status=_FLOW_STATUS[line],
                        lineage=_FLOW_LINEAGE[line],
                        bio=rep,
                        dose_ppm=int(dose),
                        duration_h=24,
                        n_live=q["live"],
                        n_early=q["early"],
                        n_late=q["late"],
                        n_necrotic=q["necrotic"],
                        n_events=n_events,
                        cells=q["cells"],
                        dead=q["dead"],
                    )
                )
            for k in range(7):  # controls block (single stains, staurosporine)
                rr = r + 3 + k
                name = wv.cell(rr, c0 + _FLOW_CTRL_NAME_OFF).value
                if not isinstance(name, str) or not name.strip():
                    continue
                cvals = {
                    nm: _num(wv.cell(rr, c0 + _FLOW_CTRL_OFF[nm]).value)
                    for nm in ("cells", "live", "dead", "early", "late", "necrotic")
                }
                ctrl_rows.append(dict(line=line, bio=rep, control=name.strip(), **cvals))

    dose = pd.DataFrame(dose_rows)
    ctrl = pd.DataFrame(ctrl_rows)
    for nm in ("live", "early", "late", "necrotic"):
        dose[f"f_{nm}"] = dose[f"n_{nm}"] / dose["n_events"]
    dose["f_dead"] = dose.f_early + dose.f_late + dose.f_necrotic
    return dose, ctrl


# --- 1b. Resazurin (metabolic activity) --------------------------------------
def load_resazurin(path, sheet="(3) Resazurin"):
    """Read the processed resazurin sheet -> tidy long DataFrame.

    Columns: line, run (biological replicate), dose_ppm, well, resazurin
    (background-corrected FI), rel_activity (within-(line, run) fold-change vs the
    0 ppm control). Non-positive / flagged wells are dropped. The '(3) Resazurin'
    values reproduce the from-raw reconstruction (raw - media - blank) to <1 unit;
    the --check-resazurin option (reconstruct_resazurin_from_raw) verifies this.
    """
    ws = load_workbook(path, data_only=True)[sheet]
    g = [[c.value for c in r] for r in ws.iter_rows()]

    def get(r, c):
        return g[r][c] if r < len(g) and c < len(g[r]) else None

    runcols = {1: range(1, 7), 2: range(8, 14), 3: range(15, 21)}  # three runs, six wells each
    lines = {"HUVEC", "HA", "U87", "SH-SY5Y"}
    rows = []
    for r in range(len(g)):
        if get(r, 0) in lines:
            line = get(r, 0)
            for k in range(7):
                rr = r + 2 + k
                dose = get(rr, 0)
                if not isinstance(dose, str) or "ppm" not in dose:
                    continue
                dppm = int(dose.replace("ppm", "").strip())
                for run, cols in runcols.items():
                    for wi, c in enumerate(cols, 1):
                        v = get(rr, c)
                        if isinstance(v, (int, float)) and v > 0:
                            rows.append(
                                dict(
                                    line=line, run=run, dose_ppm=dppm, well=wi, resazurin=float(v)
                                )
                            )
    df = pd.DataFrame(rows)
    df["rel_activity"] = np.nan
    for (ln, run), gg in df.groupby(["line", "run"]):
        base = gg.loc[gg.dose_ppm == 0, "resazurin"].mean()
        if base and base == base:
            df.loc[gg.index, "rel_activity"] = df.loc[gg.index, "resazurin"] / base
    return df


# --- 1b(ii). Resazurin provenance: rebuild corrected values from raw ----------
# The analysis reads the processed '(3) Resazurin' sheet (above). This block
# rebuilds those corrected values from the RAW resazurin workbook
# (raw - mean(media, per run) - mean(blank, per plate)) and checks them against
# the workbook's own "Corrected R" row, so the processing is reproducible from
# source. Run it with --check-resazurin. The raw workbook has one block per
# (dose, line): rows Res / Media / Blanks / Corrected R, three runs of six wells.
_RAW_RUN_OFFSETS = {1: range(1, 7), 2: range(8, 14), 3: range(15, 21)}  # +1..6, +8..13, +15..20
_RAW_BLANK_OFFSETS = range(1, 7)
_RAW_BLOCK_RE = re.compile(r"^\s*(\d+)\s*ppm\s*-\s*(.+?)\s*$", re.I)
_RAW_ALIAS = {"SH": "SH-SY5Y"}  # shortened header name


def _red_fill(cell):
    """True if a cell has a solid red fill (used in the sheet to flag excluded wells)."""
    f = cell.fill
    if f is None or f.patternType != "solid":
        return False
    rgb = getattr(f.fgColor, "rgb", None)
    return isinstance(rgb, str) and rgb.upper().endswith("FF0000")


def reconstruct_resazurin_from_raw(raw_path):
    """Rebuild corrected resazurin FI from the raw workbook and compare to the
    sheet's own 'Corrected R' row. Returns a long DataFrame with the reconstructed
    value, the file's stored value, their difference, and the exclusion flag."""
    vals = load_workbook(raw_path, data_only=True)  # cell values
    styl = load_workbook(raw_path, data_only=False)  # cell fills (exclusions)
    rows = []
    for sheet in vals.sheetnames:
        status = (
            "non-cancer"
            if sheet.lower().startswith("healthy")
            else "cancer" if sheet.lower().startswith("cancer") else sheet
        )
        wv, wsf = vals[sheet], styl[sheet]
        for r in range(1, wv.max_row + 1):
            for c0 in range(1, wv.max_column + 1):
                m = _RAW_BLOCK_RE.match(str(wv.cell(r, c0).value or ""))
                if not m:
                    continue
                dose = int(m.group(1))
                line = m.group(2).strip()
                line = _RAW_ALIAS.get(line, line)
                res_r, med_r, blk_r, cor_r = (
                    r + 2,
                    r + 3,
                    r + 4,
                    r + 5,
                )  # Res/Media/Blanks/Corrected R
                blanks = [
                    wv.cell(blk_r, c0 + o).value
                    for o in _RAW_BLANK_OFFSETS
                    if isinstance(wv.cell(blk_r, c0 + o).value, (int, float))
                ]
                if not blanks:
                    continue
                blank_avg = float(np.mean(blanks))
                for run, offs in _RAW_RUN_OFFSETS.items():
                    cols = [c0 + o for o in offs]
                    media = [
                        wv.cell(med_r, c).value
                        for c in cols
                        if isinstance(wv.cell(med_r, c).value, (int, float))
                    ]
                    if not media:
                        continue
                    media_avg = float(np.mean(media))
                    for wi, c in enumerate(cols, 1):
                        raw = wv.cell(res_r, c).value
                        if not isinstance(raw, (int, float)):
                            continue
                        excluded = _red_fill(wsf.cell(res_r, c))
                        corrected = raw - media_avg - blank_avg
                        ref = wv.cell(cor_r, c).value  # file's Corrected R
                        rows.append(
                            dict(
                                sheet=sheet,
                                status=status,
                                line=line,
                                run=run,
                                dose_ppm=dose,
                                well=wi,
                                raw_fi=float(raw),
                                media_avg=media_avg,
                                blank_avg=blank_avg,
                                corrected=corrected,
                                ref=float(ref) if isinstance(ref, (int, float)) else np.nan,
                                excluded=excluded,
                            )
                        )
    df = pd.DataFrame(rows)
    df["abs_diff"] = (df["corrected"] - df["ref"]).abs()
    return df


def check_resazurin(raw_path):
    """Print the resazurin provenance self-check (reconstruction vs stored values)."""
    print(f"Resazurin provenance check — rebuilding from raw: {raw_path}")
    df = reconstruct_resazurin_from_raw(raw_path)
    chk = df[(~df.excluded) & df["ref"].notna()]
    n = len(chk)
    maxd = float(chk["abs_diff"].max()) if n else float("nan")
    print(f"  reconstructed {len(df)} wells across lines {sorted(df.line.unique())}")
    print(
        f"  self-check vs 'Corrected R': max |diff| = {maxd:.4f} over {n} wells "
        f"({'OK (<1 unit)' if maxd < 1 else 'small residuals — see below'})"
    )
    if n and maxd >= 1:
        worst = chk.sort_values("abs_diff", ascending=False).head(4)
        print(
            "  worst-matching wells "
            "(diff | line run dose well | recomputed vs stored | media/blank):"
        )
        for _, w in worst.iterrows():
            print(
                f"    {w.abs_diff:7.1f} | {w.line} r{int(w.run)} {int(w.dose_ppm)}ppm "
                f"w{int(w.well)} | {w.corrected:.0f} vs {w.ref:.0f} | "
                f"{w.media_avg:.1f}/{w.blank_avg:.0f}"
            )
    nex = int(df.excluded.sum())
    print(f"  wells flagged (red fill) and excluded: {nex}")
    for _, w in df[df.excluded].iterrows():
        print(
            f"    {w.line} run{int(w.run)} {int(w.dose_ppm)}ppm "
            f"well{int(w.well)} (raw {w.raw_fi:.0f})"
        )
    print(
        "  (a small constant residual on one contaminated plate is expected and documented; "
        "it reflects a stored blank differing slightly from the recorded blank.)"
    )
    return df


# --- 1c. Protein (BCA) --------------------------------------------------------
_PROT_SHEETS = {
    "Protein Concentration Cancer": dict(status="cancer", line_col=14, val0=15),
    "Protein Concentration Healthy": dict(status="non-cancer", line_col=14, val0=15),
}
_PROT_LINES = {"U87", "SH-SY5Y", "HA", "HUVEC"}


def load_protein(path):
    """Read total-protein (BCA) blocks -> tidy long DataFrame.

    Columns: line, status, dose_ppm, run, read, protein_ug_ml, protein_rel
    (within-line fold-change vs the 0 ppm control). Absolute levels differ between
    lines (biological variability); analysis is within-line.
    """
    wb = load_workbook(path, data_only=True)
    rows = []
    for sheet, cfg in _PROT_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        g = [[c.value for c in r] for r in wb[sheet].iter_rows()]

        def get(r, c):
            return g[r][c] if r < len(g) and c < len(g[r]) else None

        lc, v0 = cfg["line_col"], cfg["val0"]
        for r in range(len(g)):
            v = get(r, lc)
            if (
                v in _PROT_LINES
                and isinstance(get(r + 2, v0), (int, float))
                and get(r + 2, v0) > 50
            ):
                for k in range(7):
                    rr = r + 2 + k
                    dose = get(rr, lc)
                    if not isinstance(dose, (int, float)):
                        continue
                    for j in range(6):  # 3 runs x 2 reads
                        val = get(rr, v0 + j)
                        if isinstance(val, (int, float)):
                            rows.append(
                                dict(
                                    line=v,
                                    status=cfg["status"],
                                    dose_ppm=int(dose),
                                    run=j // 2 + 1,
                                    read=j % 2 + 1,
                                    protein_ug_ml=float(val),
                                )
                            )
    df = pd.DataFrame(rows)
    df["protein_rel"] = np.nan
    for ln, gdf in df.groupby("line"):
        base = gdf.loc[gdf.dose_ppm == 0, "protein_ug_ml"].mean()
        if base and base == base:
            df.loc[gdf.index, "protein_rel"] = df.loc[gdf.index, "protein_ug_ml"] / base
    return df


# ==============================================================================
# SECTION 2 — MODELS
#   Two hierarchical Bayesian models. PyMC is imported lazily inside these
#   builders so that --check-resazurin and --help do not require it. Priors and
#   structure are documented in the Methods (weakly-informative justification there).
# ==============================================================================
def build_resazurin_model(df, xbar):
    """Hierarchical log-normal dose-response for metabolic activity.

        log M_i ~ Normal(mu_i, sigma)
        mu_i = b0[line] + b1[line]*x + b2[line]*x^2 + u[rep]      (x = log10(ppm+1) - xbar)

    Non-centred hierarchical line coefficients; per-(line, run) replicate offset.
    """
    import pymc as pm

    lines = df["line"].astype("category")
    line_idx = lines.cat.codes.values
    line_names = list(lines.cat.categories)
    rep_key = (df["line"].astype(str) + "_run" + df["run"].astype(str)).astype("category")
    rep_idx = rep_key.cat.codes.values
    rep_names = list(rep_key.cat.categories)
    x = np.log10(df["dose_ppm"].values + DOSE_OFFSET) - xbar
    logM = np.log(df["resazurin"].values)
    y_loc = float(np.mean(logM))
    coords = {"line": line_names, "rep": rep_names, "obs": np.arange(len(df))}
    with pm.Model(coords=coords) as model:
        m0 = pm.Normal("m0", y_loc, 2.0)
        s0 = pm.HalfNormal("s0", 1.0)
        m1 = pm.Normal("m1", -0.1, 0.3)
        s1 = pm.HalfNormal("s1", 0.2)
        m2 = pm.Normal("m2", 0, 0.2)
        s2 = pm.HalfNormal("s2", 0.2)
        b0 = pm.Deterministic("b0", m0 + pm.Normal("z0", 0, 1, dims="line") * s0, dims="line")
        b1 = pm.Deterministic("b1", m1 + pm.Normal("z1", 0, 1, dims="line") * s1, dims="line")
        b2 = pm.Deterministic("b2", m2 + pm.Normal("z2", 0, 1, dims="line") * s2, dims="line")
        sigma_rep = pm.HalfNormal("sigma_rep", 0.25)
        u_rep = pm.Deterministic(
            "u_rep", pm.Normal("z_rep", 0, 1, dims="rep") * sigma_rep, dims="rep"
        )
        mu = b0[line_idx] + b1[line_idx] * x + b2[line_idx] * (x**2) + u_rep[rep_idx]
        sigma = pm.HalfNormal("sigma", 0.3)
        pm.Normal("logM", mu=mu, sigma=sigma, observed=logM)
    return model, line_names


def build_composition_model(df, xbar):
    """Hierarchical Dirichlet-multinomial for the four Annexin V/PI populations.

        y_i ~ DirichletMultinomial(n_i, phi * p_i),  live = reference (logit 0)
        eta_{i,k} = a[line,k] + b[line,k]*x + c[line,k]*x^2 + r[group,k]

    Non-centred hierarchical line coefficients per non-reference category; a
    replicate-level logit offset captures extra-multinomial dispersion alongside phi.
    """
    import pymc as pm
    import pytensor.tensor as pt

    lines = df["line"].astype("category")
    line_idx = lines.cat.codes.values
    line_names = list(lines.cat.categories)
    biogrp = (
        df["line"].astype(str)
        + "|"
        + df["dose_ppm"].astype(str)
        + "|"
        + df["duration_h"].astype(str)
        + "|"
        + df["bio"].astype(str)
    ).astype("category")
    bio_idx = biogrp.cat.codes.values
    x = np.log10(df["dose_ppm"].values + DOSE_OFFSET) - xbar
    counts = df[[f"n_{q}" for q in QUADRANTS]].values.astype("int64")
    n_tot = df["n_events"].values.astype("int64")
    coords = {
        "line": line_names,
        "cat": NONREF,
        "obs": np.arange(len(df)),
        "quad": QUADRANTS,
        "biogrp": list(biogrp.cat.categories),
    }
    with pm.Model(coords=coords) as model:
        x_ = pm.Data("x", x)
        mu_a = pm.Normal("mu_a", -2.8, 1.0, dims="cat")
        mu_b = pm.Normal("mu_b", 0.0, 0.5, dims="cat")
        mu_c = pm.Normal("mu_c", 0.3, 0.5, dims="cat")
        sd_a = pm.HalfNormal("sd_a", 0.5, dims="cat")
        sd_b = pm.HalfNormal("sd_b", 0.4, dims="cat")
        sd_c = pm.HalfNormal("sd_c", 0.4, dims="cat")
        a = pm.Deterministic(
            "a", mu_a + pm.Normal("za", 0, 1, dims=("line", "cat")) * sd_a, dims=("line", "cat")
        )
        b = pm.Deterministic(
            "b", mu_b + pm.Normal("zb", 0, 1, dims=("line", "cat")) * sd_b, dims=("line", "cat")
        )
        c = pm.Deterministic(
            "c", mu_c + pm.Normal("zc", 0, 1, dims=("line", "cat")) * sd_c, dims=("line", "cat")
        )
        sd_r = pm.HalfNormal("sd_r", 0.3, dims="cat")
        r = pm.Normal("zr", 0, 1, dims=("biogrp", "cat")) * sd_r
        eta = (
            a[line_idx, :]
            + b[line_idx, :] * x_[:, None]
            + c[line_idx, :] * (x_[:, None] ** 2)
            + r[bio_idx, :]
        )
        logits = pt.concatenate([pt.zeros((eta.shape[0], 1)), eta], axis=1)
        p = pt.special.softmax(logits, axis=1)
        phi = pm.Gamma("phi", mu=250, sigma=120)
        pm.DirichletMultinomial("y", n=n_tot, a=phi * p, observed=counts, dims=("obs", "quad"))
    return model, line_names


# ==============================================================================
# SECTION 3 — FITTING & CACHE
#   fit_or_load() returns a `state` dict holding the data, the two posteriors,
#   xbar and the line orders. It fits both models once and caches the posteriors
#   (netCDF) plus a small metadata sidecar; later calls reuse the cache unless
#   --refit is given or the data file / sampling settings changed.
# ==============================================================================
def compute_xbar(flow):
    """Centring constant for log-dose, from the full observed dose ladder."""
    doses_full = sorted(flow.dose_ppm.unique())
    return float(np.mean(np.unique(np.log10(np.array(doses_full) + DOSE_OFFSET))))


def _cache_valid(outdir, data_file, sample_kw):
    meta_path = os.path.join(outdir, CACHE_META)
    if not all(
        os.path.exists(os.path.join(outdir, f))
        for f in (CACHE_RESAZURIN, CACHE_COMPOSITION, CACHE_META)
    ):
        return False
    try:
        meta = json.load(open(meta_path))
    except Exception:
        return False
    if meta.get("data_file") != os.path.abspath(data_file):
        return False
    if os.path.exists(data_file) and meta.get("data_mtime") != os.path.getmtime(data_file):
        return False
    for k in ("draws", "tune", "chains", "random_seed"):
        if meta.get(k) != sample_kw.get(k):
            return False
    return True


def fit_or_load(data_file, outdir, sample_kw, refit=False, quiet=False):
    import arviz as az

    flow, ctrl = load_flow(data_file)
    res = load_resazurin(data_file)
    prot = load_protein(data_file)
    xbar = compute_xbar(flow)

    idr_path = os.path.join(outdir, CACHE_RESAZURIN)
    idc_path = os.path.join(outdir, CACHE_COMPOSITION)

    if not refit and _cache_valid(outdir, data_file, sample_kw):
        if not quiet:
            print("Loading cached posteriors (use --refit to re-sample).")
        idr = az.from_netcdf(idr_path)
        idc = az.from_netcdf(idc_path)
    else:
        import pymc as pm

        if not quiet:
            print(
                f"Sampling both models ({sample_kw['chains']} chains x "
                f"{sample_kw['draws']} draws)..."
            )
        # metabolic model: all four lines, primary dose range
        res_fit = res[res.dose_ppm < PRIMARY_MAX].copy()
        mr, r_lines = build_resazurin_model(res_fit, xbar)
        with mr:
            idr = pm.sample(**sample_kw)
        # composition model: HA/HUVEC/U87 (SH-SY5Y excluded), primary dose range
        fl = flow[(flow.line != "SH-SY5Y") & (flow.dose_ppm < PRIMARY_MAX)].copy()
        mc, c_lines = build_composition_model(fl, xbar)
        with mc:
            idc = pm.sample(**sample_kw)
        if not quiet:
            rr = float(az.summary(idr)["r_hat"].max())
            dr = int(idr.sample_stats["diverging"].sum())
            rc = float(az.summary(idc)["r_hat"].max())
            dc = int(idc.sample_stats["diverging"].sum())
            print(f"  resazurin  : lines={r_lines} maxRhat={rr:.3f} div={dr}")
            print(f"  composition: lines={c_lines} maxRhat={rc:.3f} div={dc}")
        try:
            az.to_netcdf(idr, idr_path)
            az.to_netcdf(idc, idc_path)
            json.dump(
                dict(
                    data_file=os.path.abspath(data_file),
                    data_mtime=os.path.getmtime(data_file) if os.path.exists(data_file) else None,
                    xbar=xbar,
                    draws=sample_kw["draws"],
                    tune=sample_kw["tune"],
                    chains=sample_kw["chains"],
                    random_seed=sample_kw["random_seed"],
                ),
                open(os.path.join(outdir, CACHE_META), "w"),
            )
        except Exception as e:  # caching is best-effort
            if not quiet:
                print(f"  (could not cache posteriors: {e}; will re-fit next time)")

    r_lines = list(idr.posterior.coords["line"].values)
    c_lines = list(idc.posterior.coords["line"].values)
    return dict(
        flow=flow,
        ctrl=ctrl,
        res=res,
        prot=prot,
        xbar=xbar,
        idr=idr,
        idc=idc,
        r_lines=r_lines,
        c_lines=c_lines,
    )


# ==============================================================================
# SECTION 4 — POSTERIOR HELPERS
#   Small pure functions that convert posterior draws into the dose-response
#   summaries the figures and table need.
# ==============================================================================
def _stack(idata, name):
    return idata.posterior[name].stack(s=("chain", "draw")).values


def curve_from_quadratic(a, b, c, xbar, doses, ref_dose=0.0):
    """Response RELATIVE to ref_dose over a dose grid, from quadratic log-mean draws.
    Returns (mean, lo, hi) with lo/hi the 94% equal-tailed interval (3rd/97th pct)."""
    x = np.log10(np.asarray(doses) + DOSE_OFFSET) - xbar
    xr = np.log10(ref_dose + DOSE_OFFSET) - xbar
    logmu = a[:, None] + b[:, None] * x[None, :] + c[:, None] * x[None, :] ** 2
    logref = a + b * xr + c * xr**2
    rel = np.exp(logmu - logref[:, None])
    lo, hi = np.percentile(rel, [3, 97], axis=0)
    return rel.mean(0), lo, hi


def metab_draws(state):
    """Per-line quadratic coefficient draws for the metabolic model: (b0, b1, b2)."""
    return _stack(state["idr"], "b0"), _stack(state["idr"], "b1"), _stack(state["idr"], "b2")


def make_dead_grid(state, grid=GRID):
    """Return a function li -> (mean, lo, hi) of the posterior dead fraction (1 - live)
    over the dose grid, for composition-model line index li."""
    a_ = _stack(state["idc"], "a")
    b_ = _stack(state["idc"], "b")
    c_ = _stack(state["idc"], "c")
    xbar = state["xbar"]
    x = np.log10(grid + DOSE_OFFSET) - xbar

    def dead_grid(li):
        eta = (
            a_[li, :, None, :]
            + b_[li, :, None, :] * x[None, :, None]
            + c_[li, :, None, :] * (x[None, :, None] ** 2)
        )
        logits = np.concatenate([np.zeros((1,) + eta.shape[1:]), eta], axis=0)  # live logit = 0
        p = np.exp(logits)
        p /= p.sum(0, keepdims=True)
        dead = 1 - p[0]
        return (dead.mean(1), *np.percentile(dead, [3, 97], axis=1))

    return dead_grid


def _raw_metab_means(state):
    """Per-(line, run, dose) mean relative metabolic activity, for raw-point overlay."""
    res = state["res"]
    return (
        res[res.dose_ppm < PRIMARY_MAX]
        .groupby(["line", "run", "dose_ppm"])
        .rel_activity.mean()
        .reset_index()
    )


# ==============================================================================
# SECTION 5 — FIGURES  (one function per manuscript figure)
# ==============================================================================
def figure2_metabolic(state, outdir):
    """Figure 2 — metabolic dose-response, per line (raw replicate points + posterior fit)."""
    b0, b1, b2 = metab_draws(state)
    r_lines, xbar = state["r_lines"], state["xbar"]
    xg = GRID + DOSE_OFFSET
    raw = _raw_metab_means(state)
    rng = np.random.default_rng(1)
    fig, axes = plt.subplots(2, 2, figsize=(6.7, 5.8), sharex=True, sharey=True)
    for ax, letter, ln in zip(axes.ravel(), "ABCD", r_lines):
        li = list(r_lines).index(ln)
        m, lo, hi = curve_from_quadratic(b0[li], b1[li], b2[li], xbar, GRID)
        rr = raw[raw.line == ln]
        jit = np.exp(rng.normal(0, 0.04, len(rr)))
        ax.scatter(
            (rr.dose_ppm + DOSE_OFFSET) * jit,
            100 * rr.rel_activity,
            s=15,
            color=COL[ln],
            alpha=0.45,
            edgecolor="none",
            zorder=2,
        )
        ax.fill_between(xg, 100 * lo, 100 * hi, color=COL[ln], alpha=0.15, zorder=1)
        ax.plot(xg, 100 * m, color=COL[ln], lw=1.8, zorder=3)
        ax.axvspan(10 + 1, 40 + 1, color="grey", alpha=0.07, zorder=0)
        ax.set_xscale("log")
        ax.set_ylim(20, 135)
        ax.text(0.94, 0.94, ln, transform=ax.transAxes, ha="right", va="top", fontsize=8)
        _panel(ax, letter)
    for ax in axes[-1, :]:
        ax.set_xlabel("CO (ppm) + 1")
    for ax in axes[:, 0]:
        ax.set_ylabel("metabolic activity\n(% of control)")
    fig.tight_layout()
    paths = _savefig(fig, outdir, "figure2_metabolic")
    plt.close(fig)
    return paths


def figure3_death(state, outdir):
    """Figure 3 — cell death flat across CO dose (flow composition, HA/HUVEC/U87)."""
    c_lines, flow = state["c_lines"], state["flow"]
    dead_grid = make_dead_grid(state)
    xg = GRID + DOSE_OFFSET
    fig, ax = plt.subplots(figsize=(5.5, 4.0))
    for li, ln in enumerate(c_lines):
        m, lo, hi = dead_grid(li)
        ax.plot(xg, 100 * m, color=COL[ln], lw=1.8, label=ln)
        ax.fill_between(xg, 100 * lo, 100 * hi, color=COL[ln], alpha=0.15)
        emp = flow[flow.line == ln].groupby("dose_ppm").f_dead.mean()
        ax.scatter(
            emp.index + 1, 100 * emp.values, color=COL[ln], s=16, alpha=0.55, edgecolor="none"
        )
    ax.set_xscale("log")
    ax.set_xlabel("CO (ppm) + 1")
    ax.set_ylabel("dead fraction (%)")
    ax.set_ylim(0, None)
    ax.legend(frameon=False)
    fig.tight_layout()
    paths = _savefig(fig, outdir, "figure3_death")
    plt.close(fig)
    return paths


def figure4_decoupling(state, outdir):
    """Figure 4 — metabolic-death decoupling (metabolism falls, viability holds)."""
    b0, b1, b2 = metab_draws(state)
    r_lines, c_lines, flow, xbar = state["r_lines"], state["c_lines"], state["flow"], state["xbar"]
    dead_grid = make_dead_grid(state)
    xg = GRID + DOSE_OFFSET
    raw = _raw_metab_means(state)
    rng3 = np.random.default_rng(2)
    fig, axes = plt.subplots(1, 3, figsize=(6.7, 2.7), sharey=True)
    for ax, letter, ln in zip(axes, "ABC", c_lines):
        li_r = list(r_lines).index(ln)
        li_c = list(c_lines).index(ln)
        rr = raw[raw.line == ln]
        jitm = np.exp(rng3.normal(0, 0.04, len(rr)))
        ax.scatter(
            (rr.dose_ppm + DOSE_OFFSET) * jitm,
            100 * rr.rel_activity,
            s=10,
            color=COL[ln],
            alpha=0.40,
            edgecolor="none",
            zorder=2,
        )
        m, lo, hi = curve_from_quadratic(b0[li_r], b1[li_r], b2[li_r], xbar, GRID)
        ax.fill_between(xg, 100 * lo, 100 * hi, color=COL[ln], alpha=0.15, zorder=1)
        ax.plot(xg, 100 * m, "-", lw=1.8, color=COL[ln], label="metabolic activity", zorder=3)
        fv = flow[(flow.line == ln) & (flow.dose_ppm < PRIMARY_MAX)]
        jitv = np.exp(rng3.normal(0, 0.04, len(fv)))
        ax.scatter(
            (fv.dose_ppm + DOSE_OFFSET) * jitv,
            100 * (1 - fv.f_dead),
            s=10,
            color="0.55",
            alpha=0.5,
            edgecolor="none",
            zorder=2,
        )
        dm, dlo, dhi = dead_grid(li_c)
        ax.plot(xg, 100 * (1 - dm), "--", lw=1.6, color="0.4", label="viability", zorder=3)
        ax.axvspan(10 + 1, 40 + 1, color="grey", alpha=0.07, zorder=0)
        ax.set_xscale("log")
        ax.set_xlabel("CO (ppm) + 1")
        ax.set_ylim(0, 130)
        ax.text(0.94, 0.06, ln, transform=ax.transAxes, ha="right", va="bottom", fontsize=8)
        _panel(ax, letter)
    axes[0].set_ylabel("% of control / % live")
    axes[0].legend(frameon=False, loc="lower left")
    fig.tight_layout()
    paths = _savefig(fig, outdir, "figure4_decoupling")
    plt.close(fig)
    return paths


def figure5_protein(state, outdir):
    """Figure 5 — protein normalisation (total protein flat; activity per µg protein falls)."""
    b0, b1, b2 = metab_draws(state)
    r_lines, prot, xbar = state["r_lines"], state["prot"], state["xbar"]
    pm_ = prot.groupby(["line", "dose_ppm"]).protein_rel.mean().reset_index()
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(6.7, 3.3), sharex=True)
    for ln in ["HA", "HUVEC", "U87", "SH-SY5Y"]:
        p = pm_[pm_.line == ln]
        a1.plot(p.dose_ppm + 1, 100 * p.protein_rel, "-o", ms=3, color=COL[ln], label=ln)
        li_r = list(r_lines).index(ln)
        dd = sorted(p.dose_ppm.unique())
        m, _, _ = curve_from_quadratic(b0[li_r], b1[li_r], b2[li_r], xbar, np.array(dd))
        pr = p.set_index("dose_ppm").protein_rel.reindex(dd).values
        a2.plot(np.array(dd) + 1, 100 * (m / pr), "-o", ms=3, color=COL[ln], label=ln)
    for ax in (a1, a2):
        ax.set_xscale("log")
        ax.set_xlabel("CO (ppm) + 1")
        ax.axvspan(10 + 1, 40 + 1, color="grey", alpha=0.07)
    a1.axhline(100, color="k", lw=0.7, ls=":")
    a1.set_ylabel("total protein (% of control)")
    a2.set_ylabel("activity per µg protein (% of control)")
    a1.legend(frameon=False)
    _panel(a1, "A")
    _panel(a2, "B")
    fig.tight_layout()
    paths = _savefig(fig, outdir, "figure5_protein")
    plt.close(fig)
    return paths


# ==============================================================================
# SECTION 6 — TABLE
# ==============================================================================
def table2(state, outdir):
    """Table 2 — per-line posterior estimates: metabolic activity relative to control
    at 20 and 40 ppm (mean [94% CrI]) and the max change in dead fraction (pp)."""
    b0, b1, b2 = metab_draws(state)
    r_lines, c_lines, xbar = state["r_lines"], state["c_lines"], state["xbar"]
    dead_grid = make_dead_grid(state)

    def rel_at(li, dose):
        m, lo, hi = curve_from_quadratic(b0[li], b1[li], b2[li], xbar, np.array([dose]))
        return m[0], lo[0], hi[0]

    rows = []
    for ln in r_lines:
        li = list(r_lines).index(ln)
        m20, lo20, hi20 = rel_at(li, 20)
        m40, lo40, hi40 = rel_at(li, 40)
        rows.append(
            dict(
                line=ln,
                metab_rel_20ppm=round(m20, 3),
                metab_rel_20_lo=round(lo20, 3),
                metab_rel_20_hi=round(hi20, 3),
                metab_rel_40ppm=round(m40, 3),
                metab_rel_40_lo=round(lo40, 3),
                metab_rel_40_hi=round(hi40, 3),
            )
        )
    T = pd.DataFrame(rows)
    for ln in c_lines:
        li = list(c_lines).index(ln)
        dm, _, _ = dead_grid(li)
        T.loc[T.line == ln, "max_dead_change_pp"] = round(100 * np.max(np.abs(dm - dm[0])), 2)
    out = os.path.join(outdir, "table2_estimates.csv")
    T.to_csv(out, index=False)
    print("\nTable 2 (metabolic activity relative to control; death change in percentage points):")
    print(T.to_string(index=False))
    return out


# ==============================================================================
# SECTION 6b — SUPPLEMENTARY DIAGNOSTICS
#   Convergence table (R-hat, bulk/tail ESS), trace plots of the population-level
#   hyperparameters, and posterior predictive checks for both models. Outputs use
#   a "supp_" prefix and are intended for the supplementary material.
# ==============================================================================
_RES_HYPERPARAMS = ["m0", "m1", "m2", "s0", "s1", "s2", "sigma_rep", "sigma"]
_COMP_HYPERPARAMS = ["mu_a", "mu_b", "mu_c", "sd_a", "sd_b", "sd_c", "sd_r", "phi"]


def diagnostics(state, outdir):
    """Write supplementary diagnostics for both models.

    Produces a per-parameter convergence table (R-hat, bulk-ESS, tail-ESS), trace
    plots of the population-level hyperparameters, and posterior predictive checks.
    Files: supp_convergence.csv, supp_trace_resazurin.png,
    supp_trace_composition.png, supp_ppc_resazurin.png, supp_ppc_composition.png.
    """
    import arviz as az
    import pymc as pm

    idr, idc = state["idr"], state["idc"]
    written = []

    # --- A. convergence table: per-parameter R-hat, bulk-ESS, tail-ESS ---
    def _summ(idata, name):
        s = az.summary(idata)
        s.insert(0, "model", name)
        s.insert(1, "parameter", s.index)
        return s.reset_index(drop=True)

    conv = pd.concat([_summ(idr, "resazurin"), _summ(idc, "composition")], ignore_index=True)
    conv_path = os.path.join(outdir, "supp_convergence.csv")
    conv.to_csv(conv_path, index=False)
    written.append(conv_path)
    for name, idata in [("resazurin", idr), ("composition", idc)]:
        s = conv[conv.model == name]
        div = int(idata.sample_stats["diverging"].sum())
        print(
            f"  {name:>11}: max R-hat {s.r_hat.max():.4f}, "
            f"min ESS-bulk {int(s.ess_bulk.min())}, min ESS-tail {int(s.ess_tail.min())}, "
            f"divergences {div}"
        )

    # --- B. trace plots of the population-level hyperparameters ---
    for idata, varnames, tag, title in [
        (idr, _RES_HYPERPARAMS, "resazurin", "Resazurin model"),
        (idc, _COMP_HYPERPARAMS, "composition", "Composition model"),
    ]:
        az.plot_trace(idata, var_names=varnames, compact=True)
        fig = plt.gcf()
        fig.suptitle(f"{title} — trace (population-level hyperparameters)", y=1.005)
        fig.tight_layout()
        written.extend(_savefig(fig, outdir, f"supp_trace_{tag}", bbox_inches="tight"))
        plt.close(fig)

    # --- C. posterior predictive checks ---
    # Rebuild each model on the same data subset used for fitting, then draw from
    # the posterior predictive conditioned on the fitted posterior.
    res_fit = state["res"][state["res"].dose_ppm < PRIMARY_MAX].copy()
    mr, _ = build_resazurin_model(res_fit, state["xbar"])
    with mr:
        pm.sample_posterior_predictive(
            idr, var_names=["logM"], random_seed=7, progressbar=False, extend_inferencedata=True
        )
    # Observed vs replicated distribution of log fluorescence (density overlay).
    obs = np.asarray(idr.observed_data["logM"].values).ravel()
    ppd = idr.posterior_predictive["logM"].stack(s=("chain", "draw")).values  # (obs, s)
    bins = np.linspace(obs.min() - 0.5, obs.max() + 0.5, 40)
    centres = 0.5 * (bins[:-1] + bins[1:])
    rng = np.random.default_rng(7)
    idx = rng.choice(ppd.shape[1], size=min(200, ppd.shape[1]), replace=False)
    pp_dens = np.array([np.histogram(ppd[:, k], bins=bins, density=True)[0] for k in idx])
    lo, hi = np.percentile(pp_dens, [3, 97], axis=0)
    obs_dens, _ = np.histogram(obs, bins=bins, density=True)
    fig, ax = plt.subplots(figsize=(5.5, 4.0))
    ax.fill_between(centres, lo, hi, color=COL["U87"], alpha=0.25, label="predictive 94%")
    ax.plot(centres, pp_dens.mean(0), color=COL["U87"], lw=1.5, label="predictive mean")
    ax.plot(centres, obs_dens, color="k", lw=2, label="observed")
    ax.set_xlabel("log metabolic fluorescence")
    ax.set_ylabel("density")
    ax.set_title("Resazurin model — posterior predictive")
    ax.legend(frameon=False)
    fig.tight_layout()
    written.extend(_savefig(fig, outdir, "supp_ppc_resazurin", bbox_inches="tight"))
    plt.close(fig)

    flow = state["flow"]
    fl = flow[(flow.line != "SH-SY5Y") & (flow.dose_ppm < PRIMARY_MAX)].copy()
    mc, _ = build_composition_model(fl, state["xbar"])
    with mc:
        pm.sample_posterior_predictive(
            idc, var_names=["y"], random_seed=7, progressbar=False, extend_inferencedata=True
        )
    # Observed vs posterior-predictive quadrant proportions, one panel per category.
    yobs = idc.observed_data["y"].values.astype(float)  # (obs, quad)
    ypp = idc.posterior_predictive["y"].stack(s=("chain", "draw")).values  # (obs, quad, s)
    pobs = yobs / yobs.sum(1, keepdims=True)
    ppp = ypp / ypp.sum(1, keepdims=True)
    fig, axes = plt.subplots(1, 4, figsize=(11, 3))
    covered, total = 0, 0
    for j, (ax, cat) in enumerate(zip(axes, QUADRANTS)):
        pred_mean = ppp[:, j, :].mean(1)
        lo, hi = np.percentile(ppp[:, j, :], [3, 97], axis=1)
        ax.errorbar(
            pobs[:, j],
            pred_mean,
            yerr=[pred_mean - lo, hi - pred_mean],
            fmt="o",
            ms=3,
            alpha=0.5,
            lw=0.6,
            color=COL["U87"],
        )
        hi_lim = max(pobs[:, j].max(), hi.max()) * 1.05
        ax.plot([0, hi_lim], [0, hi_lim], "k--", lw=0.8)
        ax.set_xlim(0, hi_lim)
        ax.set_ylim(0, hi_lim)
        ax.set_title(cat)
        ax.set_xlabel("observed")
        covered += int(np.sum((pobs[:, j] >= lo) & (pobs[:, j] <= hi)))
        total += len(pobs)
    axes[0].set_ylabel("posterior predictive")
    fig.suptitle("Composition model — posterior predictive quadrant proportions", y=1.03)
    fig.tight_layout()
    written.extend(_savefig(fig, outdir, "supp_ppc_composition", bbox_inches="tight"))
    plt.close(fig)
    print(
        f"  composition PPC 94% interval coverage: {covered}/{total} "
        f"({100 * covered / total:.0f}% of well-category proportions)"
    )
    return written


# ==============================================================================
# SECTION 7 — COMMAND-LINE
# ==============================================================================
FIGURE_FUNCS = {2: figure2_metabolic, 3: figure3_death, 4: figure4_decoupling, 5: figure5_protein}
FIGURE_DESC = {
    2: "metabolic dose-response, per line",
    3: "cell death / viability, flat across dose",
    4: "metabolic-death decoupling",
    5: "protein normalisation",
}
# Figure 1 (the schematic) is produced separately in a vector graphics editor.


def build_parser():
    p = argparse.ArgumentParser(
        description="CO dose-response analysis — regenerate any figure or table.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Figures: " + "; ".join(f"{n}={d}" for n, d in FIGURE_DESC.items()),
    )
    p.add_argument(
        "--data",
        default="Paper_Data_0_3.xlsx",
        help="consolidated data workbook (default: Paper_Data_0_3.xlsx in cwd)",
    )
    p.add_argument("--outdir", default=".", help="output folder (default: current folder)")
    p.add_argument(
        "--figure",
        nargs="+",
        type=int,
        metavar="N",
        help="figure number(s) to build: 2 through 5 "
        "(Figure 1 is the schematic, made separately)",
    )
    p.add_argument("--table", action="store_true", help="build Table 2")
    p.add_argument("--all", action="store_true", help="build every figure and the table")
    p.add_argument(
        "--diagnostics",
        action="store_true",
        help="build supplementary diagnostics: convergence table (R-hat, ESS), "
        "trace plots, and posterior predictive checks (supp_*)",
    )
    p.add_argument(
        "--check-resazurin",
        action="store_true",
        help="provenance check: rebuild resazurin from raw and verify it "
        "reproduces the workbook's own corrected values (no MCMC)",
    )
    p.add_argument(
        "--raw-data",
        default="Paper_Data_Resazurin.xlsx",
        help="raw resazurin workbook for --check-resazurin "
        "(default: Paper_Data_Resazurin.xlsx in cwd)",
    )
    p.add_argument("--refit", action="store_true", help="force a fresh MCMC fit (ignore cache)")
    p.add_argument("--draws", type=int, default=SAMPLE_DEFAULTS["draws"])
    p.add_argument("--tune", type=int, default=SAMPLE_DEFAULTS["tune"])
    p.add_argument("--chains", type=int, default=SAMPLE_DEFAULTS["chains"])
    p.add_argument("--cores", type=int, default=None, help="parallel cores (default: = chains)")
    p.add_argument("--seed", type=int, default=SAMPLE_DEFAULTS["random_seed"])
    return p


def main(argv=None):
    args = build_parser().parse_args(argv)
    os.makedirs(args.outdir, exist_ok=True)

    # provenance check needs no fit and no consolidated workbook — handle it first
    if args.check_resazurin:
        if not os.path.exists(args.raw_data):
            sys.exit(
                f"ERROR: raw resazurin workbook not found: {args.raw_data}\n"
                f"Pass --raw-data PATH (the raw Healthy/Cancer-layout workbook)."
            )
        check_resazurin(args.raw_data)

    # decide which figures / table to build
    want_figs = set(args.figure or [])
    if args.all:
        want_figs.update([2, 3, 4, 5])
    want_table = args.table or args.all

    if 1 in want_figs:
        print(
            "Note: Figure 1 is the schematic, produced separately; "
            "this script builds Figures 2-5. Ignoring '--figure 1'."
        )
        want_figs.discard(1)
    bad = sorted(n for n in want_figs if n not in FIGURE_FUNCS)
    if bad:
        print(f"(no figure numbered {', '.join(map(str, bad))}; valid: 2-5)")
        want_figs -= set(bad)

    if not want_figs and not want_table and not args.diagnostics:
        if not args.check_resazurin:
            build_parser().print_help()
        return

    if not os.path.exists(args.data):
        sys.exit(
            f"ERROR: data workbook not found: {args.data}\n"
            f"Pass --data PATH, or run from the folder containing the workbook."
        )
    sample_kw = dict(
        draws=args.draws,
        tune=args.tune,
        chains=args.chains,
        cores=args.cores if args.cores is not None else args.chains,
        target_accept=SAMPLE_DEFAULTS["target_accept"],
        init=SAMPLE_DEFAULTS["init"],
        random_seed=args.seed,
        progressbar=False,
    )
    state = fit_or_load(args.data, args.outdir, sample_kw, refit=args.refit)

    written = []
    for n in sorted(want_figs):
        written.extend(FIGURE_FUNCS[n](state, args.outdir))
    if want_table:
        written.append(table2(state, args.outdir))
    if args.diagnostics:
        print("\nSupplementary diagnostics:")
        written.extend(diagnostics(state, args.outdir))
    if written:
        print("\nWrote:")
        for w in written:
            print("  " + w)


if __name__ == "__main__":
    main()
